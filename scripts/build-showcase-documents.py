"""Create the actual Word files and media archive from authored showcase content."""
import json
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = Path.cwd()
out = root / 'public/reizo/deliverables'
source = json.loads((root / 'output/showcase-artifacts-20261012/build/documents.json').read_text(encoding='utf-8'))
for key, item in source:
    doc = Document()
    for border in list(doc.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    section = doc.sections[0]
    section.page_height = Inches(11.69)
    section.page_width = Inches(8.27)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.75)
    for name in ['Normal', 'Title', 'Heading 1']:
        style = doc.styles[name]
        style.font.name = 'Microsoft YaHei'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')
    normal = doc.styles['Normal']
    normal.font.size = Pt(10)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(9)
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Heading 1'].font.size = Pt(12)
    doc.styles['Heading 1'].paragraph_format.space_before = Pt(12)
    doc.add_paragraph(item['title'], 'Title')
    for title, body in item['sections']:
        doc.add_paragraph(title, 'Heading 1')
        doc.add_paragraph(body)
    footer = section.footer.paragraphs[0]
    footer.text = 'REIZO 演示项目  /  虚构样例  /  '
    fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); footer._p.append(fld)
    footer.style = doc.styles['Normal']
    doc.save(out / f'{key}.docx')

photos = ['made-to-move','collection','sprint-sand','motion-bw','community','fabric-detail']
with ZipFile(out / 'spring-visuals.zip', 'w', ZIP_DEFLATED) as z:
    for photo in photos:
        z.write(root / f'public/reizo/assets/nynor/{photo}.jpg', photo + '.jpg')
    z.write(out / 'design-2.md', '品牌视觉规范.md')
print('Created two editable Word documents and six-image archive.')
